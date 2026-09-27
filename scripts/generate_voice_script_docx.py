import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def create_voice_recording_docx():
    doc = docx.Document()

    # Set 0.75 in margins
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Base Colors
    PRIMARY = RGBColor(27, 94, 32)      # Dark Forest Green #1B5E20
    SECONDARY = RGBColor(46, 125, 50)   # Green #2E7D32
    TEXT_DARK = RGBColor(33, 33, 33)    # Near Black
    MUTED = RGBColor(117, 117, 117)     # Grey

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("🌿 EcoBridge Mobile App")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY

    subtitle_p = doc.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(14)
    run_sub = subtitle_p.add_run("Voice Recording Script & Instructions (For Voice Artist)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = SECONDARY

    # Notice Box
    box_p = doc.add_paragraph()
    box_p.paragraph_format.space_after = Pt(14)
    r_box = box_p.add_run("Hi! Thank you so much for helping us record the audio lines for the EcoBridge app. "
                          "This app helps informal scrap collectors (waste pickers / kabadiwalas) weigh and record scrap on their phones. "
                          "Your voice will guide them step-by-step! Here is a super simple guide — no technical knowledge needed.")
    r_box.font.name = "Arial"
    r_box.font.size = Pt(11)
    r_box.font.italic = True
    r_box.font.color.rgb = TEXT_DARK

    # How to Record Section
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Super Easy Recording Instructions")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(15)
    r_h1.font.bold = True
    r_h1.font.color.rgb = PRIMARY

    instructions = [
        ("Quiet Room: ", "Record in a quiet room with fan and AC turned off and windows closed to avoid background humming."),
        ("Use Your Phone: ", "Open any normal Voice Recorder app on your phone (e.g. Samsung Voice Recorder, Apple Voice Memos, Google Recorder, or WhatsApp audio)."),
        ("Distance: ", "Hold your phone about 6 to 8 inches away from your mouth so your breath does not hit the mic."),
        ("Tone of Voice: ", "Speak in a calm, clear, friendly, and respectful tone. Imagine you are politely guiding someone who is using a smartphone for the first time."),
        ("One Clip per Line: ", "Record each line as a separate short audio clip (or record all in one go with 3-second pauses between lines)."),
        ("File Naming: ", "Please save or rename each audio file with the Name shown in the last column of the table (e.g., hi_take_photo or mr_enter_weight). Any format is okay (M4A, MP3, WAV, OGG) — our tech team will handle the rest!")
    ]

    for title, desc in instructions:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(title)
        r1.font.name = "Arial"
        r1.font.bold = True
        r1.font.size = Pt(10.5)
        r1.font.color.rgb = SECONDARY
        r2 = p.add_run(desc)
        r2.font.name = "Arial"
        r2.font.size = Pt(10.5)
        r2.font.color.rgb = TEXT_DARK

    # Function to add language table
    def add_language_section(lang_title, lang_color, lines_data):
        doc.add_page_break()
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(6)
        rh = h.add_run(lang_title)
        rh.font.name = "Arial"
        rh.font.size = Pt(15)
        rh.font.bold = True
        rh.font.color.rgb = lang_color

        # Table
        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False

        headers = ["#", "Situation / Screen", "Exact Phrase to Say (Speak Out Loud)", "Save File As"]
        col_widths = [Inches(0.4), Inches(1.8), Inches(3.6), Inches(1.4)]

        hdr_cells = table.rows[0].cells
        for i, title in enumerate(headers):
            hdr_cells[i].text = title
            hdr_cells[i].width = col_widths[i]
            set_cell_background(hdr_cells[i], "1B5E20")
            set_cell_margins(hdr_cells[i], top=120, bottom=120, left=100, right=100)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Arial"
                r.font.bold = True
                r.font.size = Pt(10)
                r.font.color.rgb = RGBColor(255, 255, 255)

        for item in lines_data:
            num, screen, phrase, english_meaning, filename = item
            row_cells = table.add_row().cells
            
            # Widths
            for i in range(4):
                row_cells[i].width = col_widths[i]
                set_cell_margins(row_cells[i], top=100, bottom=100, left=100, right=100)

            # Col 0: #
            p0 = row_cells[0].paragraphs[0]
            p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r0 = p0.add_run(f"[{num}]\n[  ]")
            r0.font.name = "Arial"
            r0.font.size = Pt(9.5)
            r0.font.bold = True

            # Col 1: Screen & context
            p1 = row_cells[1].paragraphs[0]
            r1 = p1.add_run(screen)
            r1.font.name = "Arial"
            r1.font.size = Pt(9.5)
            r1.font.bold = True
            r1.font.color.rgb = SECONDARY

            # Col 2: Exact phrase + meaning
            p2 = row_cells[2].paragraphs[0]
            r2 = p2.add_run(f"“{phrase}”\n")
            r2.font.name = "Arial"
            r2.font.bold = True
            r2.font.size = Pt(11.5)
            r2.font.color.rgb = RGBColor(10, 40, 15)

            r2_sub = p2.add_run(f"Meaning: {english_meaning}")
            r2_sub.font.name = "Arial"
            r2_sub.font.size = Pt(8.5)
            r2_sub.font.italic = True
            r2_sub.font.color.rgb = MUTED

            # Col 3: Filename
            p3 = row_cells[3].paragraphs[0]
            r3 = p3.add_run(filename)
            r3.font.name = "Consolas"
            r3.font.size = Pt(9)
            r3.font.bold = True
            r3.font.color.rgb = PRIMARY

    # Hindi Lines
    hindi_lines = [
        (1, "Home Screen", "इकोब्रिज में आपका स्वागत है। ई-कचरा रिकॉर्ड करने, भाव देखने या रीसाइक्लर खोजने के लिए नया स्क्रैप तौलें पर टैप करें।", "Welcome to EcoBridge. Tap Weigh New Scrap to start.", "hi_home_guide"),
        (2, "Camera (Step 1)", "कैमरा कबाड़ की ओर रखें और एक स्पष्ट फोटो लें।", "Point the camera at your scrap and take a clear photo.", "hi_take_photo"),
        (3, "AI Scan (Step 2)", "एआई सामग्री की पहचान कर रहा है। कृपया पहचानी गई श्रेणी की पुष्टि करें।", "AI is analyzing material. Please verify detected category.", "hi_classify_material"),
        (4, "Hazard (Battery)", "सावधान! खतरनाक बैटरी मिली है। इसे न काटें, न तोड़ें और न ही आग के पास ले जाएं।", "Warning! Dangerous battery. Do not puncture, crush or burn.", "hi_hazard_battery"),
        (5, "Hazard (CRT Glass)", "सावधान! खतरनाक सीआरटी कांच मिली है। मोटे दस्ताने पहनें और ट्यूब न तोड़ें।", "Warning! Dangerous CRT glass. Wear gloves, do not break tube.", "hi_hazard_crt"),
        (6, "Hazard (PCB)", "सावधान! सर्किट बोर्ड में लेड सोल्डर है। सीधे न छुएं और धूल या धुएं से बचें।", "Warning! Contains lead solder. Avoid bare skin contact.", "hi_hazard_pcb"),
        (7, "Weight (Step 3)", "कबाड़ का कुल वजन किलोग्राम में दर्ज करें।", "Enter total scrap weight in kilograms.", "hi_enter_weight"),
        (8, "Fair Price (Step 4)", "अपने कबाड़ का बाजार भाव और अनुमानित कुल मूल्य देखें।", "Check market rate and estimated total value.", "hi_check_price"),
        (9, "Recycler (Step 5)", "हैंडओवर के लिए अपने नजदीकी प्रमाणित रीसाइक्लर का चयन करें।", "Select a certified recycler near you for handover.", "hi_select_recycler"),
        (10, "Lot Saved", "स्क्रैप लॉट सफलतापूर्वक बन गया और आपके फोन में सुरक्षित सेव हो गया।", "Scrap lot created and saved offline on your phone.", "hi_lot_created"),
        (11, "Start Handover", "कांटे का वास्तविक वजन दर्ज करें और भुगतान राशि की पुष्टि करें।", "Enter verified scale weight and confirm payment.", "hi_start_handover"),
        (12, "Handover Done", "हस्तांतरण सफलतापूर्वक दर्ज हुआ। कृपया अपनी नकद राशि प्राप्त करें।", "Handover completed successfully. Collect your cash payment.", "hi_handover_confirmed"),
    ]

    # Marathi Lines
    marathi_lines = [
        (1, "Home Screen", "इकोब्रिजमध्ये आपले स्वागत आहे. ई-कचरा नोंदवण्यासाठी, भाव पाहण्यासाठी किंवा रिसायकलर शोधण्यासाठी नवीन स्क्रॅप मोजा वर टॅप करा.", "Welcome to EcoBridge. Tap Weigh New Scrap to start.", "mr_home_guide"),
        (2, "Camera (Step 1)", "कॅमेरा भंगाराकडे धरा आणि एक स्पष्ट फोटो काढा.", "Point the camera at your scrap and take a clear photo.", "mr_take_photo"),
        (3, "AI Scan (Step 2)", "एआय भंगाराची तपासणी करत आहे. कृपया ओळखलेल्या प्रकाराची खात्री करा.", "AI is analyzing material. Please verify detected category.", "mr_classify_material"),
        (4, "Hazard (Battery)", "सावधान! धोकादायक बॅटरी आढळली आहे. याला कापू नका, तोडू नका किंवा आगीजवळ नेऊ नका.", "Warning! Dangerous battery. Do not puncture, crush or burn.", "mr_hazard_battery"),
        (5, "Hazard (CRT Glass)", "सावधान! धोकादायक सीआरटी काच आढळली आहे. जाड हातमोजे वापरा आणि काच फोडू नका.", "Warning! Dangerous CRT glass. Wear gloves, do not break tube.", "mr_hazard_crt"),
        (6, "Hazard (PCB)", "सावधान! सर्किट बोर्डमध्ये लेड सोल्डर आहे. थेट हात लावू नका आणि विषारी धुरापासून दूर राहा.", "Warning! Contains lead solder. Avoid bare skin contact.", "mr_hazard_pcb"),
        (7, "Weight (Step 3)", "भंगाराचे एकूण वजन किलोग्रॅममध्ये टाका.", "Enter total scrap weight in kilograms.", "mr_enter_weight"),
        (8, "Fair Price (Step 4)", "आपल्या भंगाराचा बाजारभाव आणि अंदाजित एकूण मूल्य तपासा.", "Check market rate and estimated total value.", "mr_check_price"),
        (9, "Recycler (Step 5)", "हस्तांतरणासाठी जवळचा अधिकृत रिसायकलर निवडा.", "Select a certified recycler near you for handover.", "mr_select_recycler"),
        (10, "Lot Saved", "स्क्रॅप लॉट यशस्वीरीत्या तयार झाला आणि फोनमध्ये ऑफलाइन सेव्ह झाला.", "Scrap lot created and saved offline on your phone.", "mr_lot_created"),
        (11, "Start Handover", "काट्यावरील प्रत्यक्ष वजन नोंदवा आणि पेमेंट रकमेची खात्री करा.", "Enter verified scale weight and confirm payment.", "mr_start_handover"),
        (12, "Handover Done", "हस्तांतरण यशस्वीरीत्या नोंदवले गेले. कृपया आपली रोख रक्कम स्वीकारा.", "Handover completed successfully. Collect your cash payment.", "mr_handover_confirmed"),
    ]

    # English Lines
    english_lines = [
        (1, "Home Screen", "Welcome to EcoBridge. Tap Weigh New Scrap to record lots, check prices, or find recyclers.", "Dashboard welcome prompt", "en_home_guide"),
        (2, "Camera (Step 1)", "Point the camera at your scrap lot and take a clear photo.", "Photo capture instruction", "en_take_photo"),
        (3, "AI Scan (Step 2)", "AI is analyzing the scrap material. Please verify the detected category.", "AI scan confirmation prompt", "en_classify_material"),
        (4, "Hazard (Battery)", "Warning! Hazardous battery detected. Do not puncture, crush, or expose to heat.", "Battery PPE & safety alert", "en_hazard_battery"),
        (5, "Hazard (CRT Glass)", "Warning! Hazardous CRT glass detected. Wear protective gloves and do not break the tube.", "CRT vacuum implosion alert", "en_hazard_crt"),
        (6, "Hazard (PCB)", "Warning! Contains lead solder. Avoid bare skin contact and do not inhale dust.", "PCB lead solder warning", "en_hazard_pcb"),
        (7, "Weight (Step 3)", "Enter or adjust the gross weight of the scrap lot in kilograms.", "Weight keypad prompt", "en_enter_weight"),
        (8, "Fair Price (Step 4)", "Review the benchmark market price and estimated total value for your lot.", "Price intelligence prompt", "en_check_price"),
        (9, "Recycler (Step 5)", "Select a certified recycler near you to proceed with handover.", "Recycler selection prompt", "en_select_recycler"),
        (10, "Lot Saved", "Scrap lot created successfully and saved offline on your device.", "Offline lot creation confirmation", "en_lot_created"),
        (11, "Start Handover", "Enter the verified scale weight and confirm the payment amount.", "Handover scale terminal prompt", "en_start_handover"),
        (12, "Handover Done", "Handover recorded successfully. Please collect your cash payment.", "Cash settlement confirmation", "en_handover_confirmed"),
    ]

    add_language_section("2. Hindi Voice Lines (हिंदी में क्या बोलना है)", PRIMARY, hindi_lines)
    add_language_section("3. Marathi Voice Lines (मराठीत काय बोलायचे आहे)", PRIMARY, marathi_lines)
    add_language_section("4. English Voice Lines (Optional Reference)", PRIMARY, english_lines)

    # Final section: Checklist & Hand-off
    doc.add_page_break()
    h_end = doc.add_paragraph()
    h_end.paragraph_format.space_before = Pt(8)
    h_end.paragraph_format.space_after = Pt(6)
    r_end = h_end.add_run("5. How to Send the Audio Files Back")
    r_end.font.name = "Arial"
    r_end.font.size = Pt(15)
    r_end.font.bold = True
    r_end.font.color.rgb = PRIMARY

    p_end = doc.add_paragraph()
    p_end.paragraph_format.space_after = Pt(8)
    r_pe = p_end.add_run(
        "Once you finish recording:\n\n"
        "1. Put all your audio files into a single folder on your phone or laptop.\n"
        "2. You can send them via Google Drive link, WhatsApp (as Document to avoid compression), or email.\n"
        "3. Don't worry about converting to .ogg — our engineering team will optimize and convert them automatically!\n\n"
        "Thank you so much for your time and voice! You are helping empower thousands of informal waste pickers across India! 🇮🇳"
    )
    r_pe.font.name = "Arial"
    r_pe.font.size = Pt(11)
    r_pe.font.color.rgb = TEXT_DARK

    out_path = "c:/Users/admin/Downloads/Ecobridge-repo/EcoBridge_Voice_Recording_Script_For_Voice_Artist.docx"
    doc.save(out_path)
    print(f"Successfully generated Word document at: {out_path}")

if __name__ == "__main__":
    create_voice_recording_docx()
